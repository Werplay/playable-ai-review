from __future__ import annotations

import json
from pathlib import Path


def verdict(findings):
    confirmed = [f for f in findings if f.classification == "verified"]
    uncertain = [f for f in findings if f.classification == "needs_human_review"]
    if any(f.severity == "critical" for f in confirmed): return "critical risk"
    if any(f.severity == "high" for f in confirmed): return "high risk"
    if uncertain:
        highest = "critical" if any(f.severity == "critical" for f in uncertain) else "high" if any(f.severity == "high" for f in uncertain) else "non-high"
        return f"human review required ({highest} candidate)"
    return "no concrete introduced defects found"


def write_reports(out: Path, data: dict) -> list[Path]:
    out.mkdir(parents=True, exist_ok=True)
    json_path = out / "report.json"; json_path.write_text(json.dumps(data, indent=2)+"\n")
    findings = data["findings"]
    lines = ["# Local AI code review", "", "## Summary", "", f"**Verdict:** {data['verdict']}", f"**Coverage:** {data['coverage']['reviewed']}/{data['coverage']['eligible']} eligible changed files; {'complete' if data['coverage']['complete'] else 'incomplete'}", f"**Index:** {data['index']['files']} tracked text files, {data['index']['chunks']} chunks ({data['index']['embedding_model']}, {data['index']['dimensions']} dimensions)", ""]
    for classification, heading in (("verified","Confirmed findings"),("needs_human_review","Needs human review")):
        lines += [f"## {heading}", ""]
        selected=[f for f in findings if f["classification"]==classification]
        if not selected: lines += ["None.", ""]
        for f in selected:
            lines += [f"### {f['severity'].upper()}: {f['category']} — `{f['file']}:{f['line']}`", "", f["explanation"], "", f"Evidence: `{f['evidence']}`", "", f"Scenario: {f['scenario']}", "", f"Recommendation: {f['recommendation']}", ""]
    lines += ["## Rejected candidates", "", f"{data['rejected_count']} rejected; details are retained in terminal diagnostics only, not published here.", "", "## Quality and security heuristics", "", "Reviewed correctness boundaries, async/error behavior, authorization and isolation, input/output security, cache and time-unit scoping, collection semantics, cleanup, retries, and client/server contracts.", "", "## Testing recommendations", "", "Exercise the confirmed failure scenarios above with the repository's normal test tooling. No project code was executed by this reviewer.", "", "## RAG retrieval summary", ""]
    lines += [f"- `{x['path']}:{x['start_line']}-{x['end_line']}`" for x in data.get("retrieval",[])] or ["No context chunks retrieved."]
    md_path=out/"report.md"; md_path.write_text("\n".join(lines)+"\n")
    paths=[md_path,json_path]
    try:
        from docx import Document
        doc=Document(); doc.add_heading("Local AI code review",0)
        for line in lines:
            if line.startswith("### "): doc.add_heading(line[4:],2)
            elif line.startswith("## "): doc.add_heading(line[3:],1)
            elif line and not line.startswith("#"): doc.add_paragraph(line.replace("**", "").replace("`", ""))
        docx=out/"report.docx"; doc.save(docx); paths.append(docx)
    except ImportError: pass
    return paths
